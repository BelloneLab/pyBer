"""Verify inferred tables remain usable through GUI selection and project saves."""
from pathlib import Path
import copy
import tempfile
import json
import unittest
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
import pandas as pd

import test_postprocessing_empty_state as gui_fixture
from gui_postprocessing import _load_behavior_csv, _load_behavior_ethovision, _behavior_table_info
from analysis_core import ProcessedTrial
from PySide6 import QtCore, QtTest, QtWidgets


class BehaviorImportGuiTests(unittest.TestCase):
    """Use isolated preferences and real widgets without touching user settings."""

    setUpClass = classmethod(gui_fixture.PostprocessingEmptyStateTests.setUpClass.__func__)
    setUp = gui_fixture.PostprocessingEmptyStateTests.setUp
    tearDown = gui_fixture.PostprocessingEmptyStateTests.tearDown

    def test_comparison_legend_is_outside_data_and_scrolls_without_losing_labels(self):
        view = self.panel.behavior_zone_panel
        self.panel.combo_event_kind.setCurrentText("Behavior")
        rows = [dict(file_id="test", label=f"behavior <{i}> with a long name", events=2,
                     before=1., during=2., after=3., during_minus_before=1., after_minus_before=2.)
                for i in range(24)]
        view.scope.setCurrentIndex(0)
        view._show_results(rows, [])
        self.panel.resize(1000, 900)
        self.panel.show()
        for _ in range(5):
            self.app.processEvents()
        self.assertIsNone(view.response.plotItem.legend)
        self.assertEqual(view.response_legend_layout.count(), 25)
        for i, row in enumerate(rows):
            labels = view.response_legend_layout.itemAt(i).widget().findChildren(QtWidgets.QLabel)
            self.assertEqual(labels[-1].text(), row["label"])
            self.assertEqual(labels[-1].textFormat(), QtCore.Qt.TextFormat.PlainText)
        self.assertLessEqual(view.response_legend.geometry().bottom(), view.response.geometry().top())
        self.assertGreater(view.response_legend.horizontalScrollBar().maximum(), 0)
        view._clear_results()
        self.assertEqual(view.response_legend_layout.count(), 0)
        self.assertTrue(view.response_legend.isHidden())

    def test_clock_switch_spatial_selection_and_project_roundtrip(self):
        """Changing clocks updates derived events and survives an HDF5 roundtrip."""
        table = pd.DataFrame({
            "timestamp_software": [100., 101., 102., 103.],
            "timestamp_camera": [50., 51., 52., 53.],
            "behavior_rearing": [0, 1, 1, 0],
            "mouse_1_center_x": [10., 11., 12., -1.],
            "mouse_1_center_y": [20., 21., 22., -1.],
            "mouse_2_center_x": [30., 31., 32., 33.],
            "mouse_2_center_y": [40., 41., 42., 43.],
            "live_detection_available": [1, 1, 1, 0],
        })
        with tempfile.TemporaryDirectory() as directory:
            csv_path = Path(directory) / "recording_metadata.csv"
            table.to_csv(csv_path, index=False)
            info = _load_behavior_csv(str(csv_path))
            panel = self.panel
            panel._behavior_sources = {"recording": info}
            panel._update_behavior_time_panel()
            panel._refresh_spatial_columns()
            self.assertEqual(panel.combo_spatial_x.currentText(), "mouse_1_center_x")
            self.assertEqual(panel.combo_spatial_y.currentText(), "mouse_1_center_y")
            self.assertEqual(set(info["behaviors"]), {"behavior_rearing"})
            np.testing.assert_equal(info["time"], [100, 101, 102, 103])
            self.assertTrue(np.isnan(info["trajectory"]["mouse_1_center_x"][-1]))
            info["event_behaviors"] = {"position": {
                "variable": "mouse_1_center_x", "rule": "> 10.5 and < 11.5", "align": "Align to onset",
            }}
            panel.combo_behavior_clock.setCurrentText("timestamp_camera")
            np.testing.assert_equal(info["time"], [50, 51, 52, 53])
            np.testing.assert_equal(info["event_behaviors"]["position"]["on"], [51])
            panel.combo_spatial_x.setCurrentText("mouse_2_center_x")
            panel.combo_spatial_y.setCurrentText("mouse_2_center_y")
            panel._refresh_spatial_columns()
            self.assertEqual(panel.combo_spatial_x.currentText(), "mouse_2_center_x")
            self.assertEqual(panel.combo_spatial_y.currentText(), "mouse_2_center_y")
            project_path = str(Path(directory) / "roundtrip.h5")
            panel._save_project_h5(project_path)
            saved = panel._load_project_h5(project_path)
            loaded = saved["behavior_sources"]["recording"]
            np.testing.assert_equal(loaded["time"], info["time"])
            for clock, values in info["time_candidates"].items():
                np.testing.assert_equal(loaded["time_candidates"][clock], values)
            self.assertEqual(loaded["auto_time_column"], "timestamp_software")
            self.assertEqual(loaded["import_report"], json.loads(json.dumps(info["import_report"])))
            panel._behavior_sources = {"recording": loaded}
            panel.combo_behavior_clock.setCurrentText("Auto")
            np.testing.assert_equal(loaded["time"], [100, 101, 102, 103])
            np.testing.assert_equal(loaded["event_behaviors"]["position"]["on"], [101])

    def test_generic_time_generation_and_explicit_event_timestamps(self):
        """FPS fallback and explicit event lists keep their existing semantics."""
        info = _behavior_table_info(pd.DataFrame({"groom": [0, 1, 0]}), "binary_columns", 2.)
        self.assertTrue(info["needs_generated_time"])
        np.testing.assert_equal(info["time"], [0, .5, 1])
        info = _behavior_table_info(pd.DataFrame({"timestamp_software": [100, 101, 102],
                                                "groom": [12., np.nan, 15.]}), "timestamp_columns", 30.)
        self.assertEqual(set(info["behaviors"]), {"groom"})
        np.testing.assert_equal(info["behaviors"]["groom"], [12, 15])

    def test_mamir_behavior_zone_pairing_and_project_roundtrip(self):
        """One chosen animal stays paired with its fiber file across a save."""
        panel = self.panel
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fiber = root / "PreTest_day1_8534_0001_AIN01.csv"
            panel._processed = [ProcessedTrial(
                path=str(fiber), channel_id="AIN01", time=np.arange(10., 20., .1),
                raw_signal=np.zeros(100), raw_reference=np.zeros(100),
                output=np.arange(10., 20., .1), output_label="dFF")]
            frames = root / "behavior_frames.csv"
            pd.DataFrame([
                (0, 12., 1, 0), (0, 12., 2, 0),
                (1, 13., 1, 0), (1, 13., 2, 1),
                (2, 14., 1, 1), (2, 14., 2, 1),
                (3, 15., 1, 0), (3, 15., 2, 0),
            ], columns=["frame", "time_s", "identity", "attack"]).to_csv(frames, index=False)
            with patch.object(QtWidgets.QFileDialog, "getOpenFileNames", return_value=([str(frames)], "")), \
                 patch.object(QtWidgets.QInputDialog, "getItem",
                              side_effect=[("2", True), (f"1. {fiber}", True)]):
                panel._load_behavior_zone_files("behavior")
            key = fiber.stem
            np.testing.assert_allclose(panel._behavior_sources[key]["event_behaviors"]["attack"]["on"], [13.])
            panel.combo_behavior_name.setCurrentText("attack")
            on, duration = panel._get_events_for_proc(panel._processed[0])
            np.testing.assert_allclose(on, [13.])
            np.testing.assert_allclose(duration, [2.])

            pd.DataFrame([(2, "target", 1, 120, 129, 10, 1.)],
                         columns=["animal_id", "zone", "bout", "start_frame",
                                  "end_frame", "observed_frames", "duration_s"]
                         ).to_csv(root / "zone_bouts.csv", index=False)
            (root / "summary.json").write_text(json.dumps({"fps": 10}), encoding="utf-8")
            with patch.object(QtWidgets.QFileDialog, "getOpenFileNames",
                              return_value=([str(root / "zone_bouts.csv")], "")), \
                 patch.object(QtWidgets.QInputDialog, "getItem", return_value=(f"1. {fiber}", True)):
                panel._load_behavior_zone_files("behavior")
            self.assertIn("Zone: target", panel._behavior_sources[key]["event_behaviors"])
            other = ProcessedTrial(path=str(root / "unpaired_AIN01.csv"), channel_id="AIN01",
                                   time=np.arange(10., 20., .1), raw_signal=np.zeros(100),
                                   raw_reference=np.zeros(100), output=np.zeros(100), output_label="dFF")
            self.assertIsNone(panel._match_behavior_source(other))

            explorer = panel.behavior_zone_panel
            panel._open_behavior_zone_explorer()
            self.assertFalse(explorer.isHidden())
            behavior_items = explorer.selectors["behavior"]
            self.assertEqual(behavior_items.count(), 1)
            behavior_items.item(0).setCheckState(QtCore.Qt.CheckState.Checked)
            explorer._compute()
            self.assertEqual(explorer._last_rows[0]["events"], 1)
            self.assertEqual(explorer.table.rowCount(), 1)
            panel._set_event_category("zone")
            self.assertTrue(explorer.isHidden())
            panel.combo_behavior_name.setCurrentText("Zone: target")
            on, duration = panel._get_events_for_proc(panel._processed[0])
            np.testing.assert_allclose(on, [12.])
            np.testing.assert_allclose(duration, [1.])

            # Group mode must use one independently paired source per fiber file.
            panel._processed.append(other)
            panel._behavior_sources[Path(other.path).stem] = copy.deepcopy(panel._behavior_sources[key])
            panel._behavior_sources[Path(other.path).stem]["import_report"].update(
                paired_index=1, paired_path=other.path)
            explorer.reload()
            panel._set_event_category("behavior")
            panel.tab_visual_mode.setCurrentIndex(1)
            explorer._compute()
            self.assertEqual(explorer._last_group_rows[0]["recordings"], 2)
            self.assertEqual(explorer._last_group_rows[0]["events"], 2)

            project = str(root / "mamir_project.h5")
            panel._save_project_h5(project)
            restored = panel._load_project_h5(project)["behavior_sources"][key]
            np.testing.assert_allclose(restored["event_behaviors"]["attack"]["on"], [13.])
            np.testing.assert_allclose(restored["event_behaviors"]["Zone: target"]["off"], [13.])

    def test_project_open_restores_embedded_behavior_without_source_files(self):
        """Both project Open and signal/drop Open restore the complete snapshot."""
        panel = self.panel
        with tempfile.TemporaryDirectory() as directory:
            time = np.arange(0., 60., .1)
            source_path = Path(directory) / "recording_metadata.csv"
            pd.DataFrame({
                "timestamp_software": time,
                "timestamp_camera": time + .25,
                "rearing": ((time % 10 >= 4) & (time % 10 < 6)).astype(int),
                "mouse_center_x": np.sin(time), "mouse_center_y": np.cos(time),
            }).to_csv(source_path, index=False)
            panel._processed = [ProcessedTrial(
                path=str(Path(directory) / "recording.csv"), channel_id="AIN01",
                time=time, raw_signal=np.sin(time), raw_reference=np.cos(time),
                output=np.sin(time), output_label="dFF")]
            panel._load_behavior_paths([str(source_path)], replace=True)
            panel._refresh_behavior_list()
            panel.combo_behavior_clock.setCurrentText("timestamp_camera")
            panel.combo_behavior_name.setCurrentText("rearing")
            panel.spin_b0.setValue(-3.)
            panel.spin_b1.setValue(-1.)
            expected = panel._behavior_sources["recording_metadata"]
            project = str(Path(directory) / "snapshot.h5")
            panel._save_project_h5(project)
            # Existing linked files must not cause a reload prompt or override.
            with patch.object(QtWidgets.QMessageBox, "question") as question:
                self.assertTrue(panel._load_project_from_path(project))
                question.assert_not_called()
            source_path.unlink()
            for route in ("project", "signal", "drop"):
                with self.subTest(route=route):
                    panel._behavior_sources = {}
                    panel._processed = []
                    panel.combo_behavior_clock.setCurrentText("Auto")
                    with patch.object(QtWidgets.QMessageBox, "question") as question:
                        if route == "project":
                            self.assertTrue(panel._load_project_from_path(project))
                        elif route == "signal":
                            with patch.object(QtWidgets.QFileDialog, "getOpenFileNames", return_value=([project], "")):
                                panel.btn_load_processed_single.click()
                        else:
                            panel._on_preprocessed_files_dropped([project])
                        question.assert_not_called()
                    restored = panel._behavior_sources["recording_metadata"]
                    self.assertEqual(len(panel._processed), 1)
                    self.assertEqual(panel.combo_behavior_name.currentText(), "rearing")
                    self.assertEqual(panel.combo_behavior_clock.currentText(), "timestamp_camera")
                    self.assertIn("1 file(s) loaded", panel.lbl_beh.text())
                    np.testing.assert_equal(restored["time"], expected["time"])
                    np.testing.assert_equal(restored["behaviors"]["rearing"], expected["behaviors"]["rearing"])
                    for name, values in expected["trajectory"].items():
                        np.testing.assert_equal(restored["trajectory"][name], values)
                    for name, values in expected["time_candidates"].items():
                        np.testing.assert_equal(restored["time_candidates"][name], values)
                    self.assertGreater(np.asarray(panel._last_events).size, 0)

    def test_zone_workbook_sheet_choice_and_explicit_fiber_pairing(self):
        """Arena selection feeds the existing PSTH without crossing fiber pairings."""
        panel = self.panel
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fiber = root / "fiber_8534.csv"
            other = root / "fiber_9999.csv"
            panel._processed = [ProcessedTrial(
                path=str(path), channel_id="AIN01", time=np.arange(10., 20., .1),
                raw_signal=np.zeros(100), raw_reference=np.zeros(100),
                output=np.arange(100.), output_label="dFF") for path in (fiber, other)]
            workbook = root / "zones.xlsx"
            with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
                pd.DataFrame({"Trial time": [12, 13, 14, 15],
                              "In zone": [0, 0, 0, 0]}).to_excel(writer, sheet_name="Sheet1", index=False)
                pd.DataFrame({"Trial time": [12, 13, 14, 15],
                              "In zone": [0, 1, 1, 0]}).to_excel(writer, sheet_name="Sheet2", index=False)
            with patch.object(QtWidgets.QFileDialog, "getOpenFileNames",
                              return_value=([str(workbook)], "")), \
                 patch.object(QtWidgets.QInputDialog, "getItem",
                              return_value=(f"1. {fiber}", True)):
                panel._load_behavior_zone_files("zone")
                explorer = panel.behavior_zone_panel
                self.assertNotIn(fiber.stem, panel._behavior_sources)
                self.assertEqual(explorer.arena_combo.count(), 3)
                self.assertEqual(explorer.arena_combo.currentData(), "")
                explorer.arena_combo.setCurrentIndex(explorer.arena_combo.findData("Sheet2"))
                explorer.arena_combo.activated.emit(explorer.arena_combo.currentIndex())
            self.assertEqual(panel._behavior_sources[fiber.stem]["sheet"], "Sheet2")
            self.assertIn("Zone: In zone", panel._behavior_sources[fiber.stem]["behaviors"])
            self.assertIsNone(panel._match_behavior_source(panel._processed[1]))
            self.assertIs(explorer.arena_combo, panel.combo_arena)
            self.assertTrue(explorer.isHidden())
            self.assertIn("Sheet2", explorer.pairing.text())
            panel.combo_behavior_name.setCurrentText("Zone: In zone")
            on, duration = panel._get_events_for_proc(panel._processed[0])
            np.testing.assert_allclose(on, [13.])
            np.testing.assert_allclose(duration, [2.])
            panel._behavior_sources[fiber.stem]["event_behaviors"]["attack"] = {
                "on": np.array([13.]), "off": np.array([14.]), "dur": np.array([1.])}
            explorer.arena_combo.setCurrentIndex(explorer.arena_combo.findData("Sheet1"))
            explorer.arena_combo.activated.emit(explorer.arena_combo.currentIndex())
            self.assertEqual(panel._behavior_sources[fiber.stem]["sheet"], "Sheet1")
            self.assertIn("attack", panel._behavior_sources[fiber.stem]["event_behaviors"])
            np.testing.assert_array_equal(panel._behavior_sources[fiber.stem]["behaviors"]["Zone: In zone"],
                                          [0, 0, 0, 0])
            explorer.arena_combo.setCurrentIndex(explorer.arena_combo.findData("Sheet2"))
            explorer.arena_combo.activated.emit(explorer.arena_combo.currentIndex())
            project = str(root / "zones_project.h5")
            panel._save_project_h5(project)
            restored = panel._load_project_h5(project)["behavior_sources"][fiber.stem]
            self.assertEqual(restored["sheet"], "Sheet2")
            self.assertEqual(len(restored["import_report"]["arena_options"]), 2)
            self.assertIn("attack", restored["event_behaviors"])
            np.testing.assert_array_equal(restored["behaviors"]["Zone: In zone"], [0, 1, 1, 0])

    def test_ethovision_arena_metadata_is_shown_before_import(self):
        """Multiple tracking sheets are labeled with arena and animal metadata."""
        import openpyxl

        with tempfile.TemporaryDirectory() as directory:
            workbook = openpyxl.Workbook()
            first = workbook.active
            first.title = "Track-Arena 1-Subject 1"
            second = workbook.create_sheet("Track-Arena 2-Subject 1")
            for sheet, arena, subject in ((first, "Arena 1", "8534 R"),
                                          (second, "Arena 2", "8534 rien")):
                sheet.append(["Arena name", arena])
                sheet.append(["<User-defined 1>", subject])
                sheet.append(["Trial time", "In zone"])
                sheet.append(["s", ""])
                sheet.append([0.0, 1])
            path = str(Path(directory) / "arena.xlsx")
            workbook.save(path)
            options = self.panel._inspect_arena_workbook(path)
            self.assertEqual([item["label"] for item in options],
                             ["Arena 1 — 8534 R", "Arena 2 — 8534 rien"])

    def test_auto_ethovision_keeps_legacy_interpolation_and_events(self):
        """Missing tracking samples must be cleaned exactly as in main's loader."""
        import openpyxl

        with tempfile.TemporaryDirectory() as directory:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.append(["Arena name", "Arena 1"])
            sheet.append(["Trial time", "In zone", "X center", "Y center"])
            sheet.append(["s", "", "cm", "cm"])
            for row in [(0, 0, 10, 20), (1, 1, 11, 21), (2, None, None, None),
                        (3, 1, 13, 23), (4, 0, 14, 24)]:
                sheet.append(row)
            path = Path(directory) / "legacy_zones.xlsx"
            workbook.save(path)
            original_bytes = path.read_bytes()
            legacy = _load_behavior_ethovision(str(path), sheet_name=sheet.title)
            self.panel._add_generic_behavior_zone_file(str(path), "zone", sheet_name=sheet.title)
            imported = next(iter(self.panel._behavior_sources.values()))
            np.testing.assert_array_equal(imported["time"], legacy["time"])
            np.testing.assert_array_equal(imported["behaviors"]["Zone: In zone"],
                                          legacy["behaviors"]["In zone"])
            np.testing.assert_array_equal(imported["behaviors"]["Zone: In zone"], [0, 1, 1, 1, 0])
            for name, values in legacy["trajectory"].items():
                np.testing.assert_array_equal(imported["trajectory"][name], values)
            for expected, actual in zip(self.panel._extract_behavior_events(legacy, "In zone"),
                                        self.panel._extract_behavior_events(imported, "Zone: In zone")):
                np.testing.assert_array_equal(actual, expected)
            self.assertEqual(path.read_bytes(), original_bytes)

    def test_setup_load_uses_auto_detection_without_replacing_existing_sources(self):
        """The familiar Setup action leaves zone analysis in the original dashboard."""
        panel = self.panel
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fiber = root / "fiber_8534.csv"
            panel._processed = [ProcessedTrial(
                path=str(fiber), channel_id="AIN01", time=np.arange(10., 20., .1),
                raw_signal=np.zeros(100), raw_reference=np.zeros(100),
                output=np.arange(100.), output_label="dFF")]
            old = _behavior_table_info(pd.DataFrame({"Trial time": [10., 11., 12.],
                                                      "groom": [0, 1, 0]}), "binary_columns", 0.)
            panel._behavior_sources["unrelated_recording"] = old
            zones = root / "arena_zones.csv"
            pd.DataFrame({"Trial time": [12., 13., 14., 15.],
                          "In zone": [0, 1, 1, 0]}).to_csv(zones, index=False)
            with patch.object(QtWidgets.QFileDialog, "getOpenFileNames",
                              return_value=([str(zones)], "")):
                panel._load_behavior_files()
            self.assertIs(panel._behavior_sources["unrelated_recording"], old)
            self.assertIn("Zone: In zone", panel._behavior_sources[fiber.stem]["behaviors"])
            self.assertEqual(panel._event_category(), "zone")
            self.assertTrue(panel.behavior_zone_panel.isHidden())

    def test_setup_timestamp_mode_keeps_legacy_loader(self):
        panel = self.panel
        panel.combo_behavior_file_type.setCurrentIndex(1)
        with patch.object(QtWidgets.QFileDialog, "getOpenFileNames",
                          return_value=(["event_timestamps.csv"], "")), \
             patch.object(panel, "_load_behavior_paths") as legacy:
            panel._load_behavior_files()
        legacy.assert_called_once_with(["event_timestamps.csv"], replace=False)

    def test_setup_timestamp_mode_never_overwrites_existing_source(self):
        panel = self.panel
        panel.combo_behavior_file_type.setCurrentIndex(1)
        existing = {"kind": "timestamp_columns", "behaviors": {"attack": np.array([12.])}}
        panel._behavior_sources["event_timestamps"] = existing
        with patch.object(QtWidgets.QFileDialog, "getOpenFileNames",
                          return_value=(["event_timestamps.csv"], "")), \
             patch.object(QtWidgets.QMessageBox, "warning") as warning, \
             patch.object(panel, "_load_behavior_paths") as legacy:
            panel._load_behavior_files()
        legacy.assert_not_called()
        warning.assert_called_once()
        self.assertIs(panel._behavior_sources["event_timestamps"], existing)

    def test_mixed_arena_sheet_keeps_selected_zone_mode(self):
        """Arena changes preserve Zone mode when a sheet contains both kinds."""
        panel = self.panel
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            fiber = root / "fiber.csv"
            panel._processed = [ProcessedTrial(
                path=str(fiber), channel_id="AIN01", time=np.arange(10., 20., .1),
                raw_signal=np.zeros(100), raw_reference=np.zeros(100),
                output=np.arange(100.), output_label="dFF")]
            book = root / "mixed.xlsx"
            with pd.ExcelWriter(book, engine="openpyxl") as writer:
                for sheet in ("Arena 1", "Arena 2"):
                    pd.DataFrame({"Trial time": [12., 13., 14., 15.],
                                  "Moving": [0, 1, 1, 0],
                                  "In zone": [0, 1, 1, 0]}).to_excel(writer, sheet_name=sheet, index=False)
            explorer = panel.behavior_zone_panel
            panel._set_event_category("zone")
            with patch.object(QtWidgets.QFileDialog, "getOpenFileNames",
                              return_value=([str(book)], "")):
                panel._load_behavior_zone_files("behavior")
            index = explorer.arena_combo.findData("Arena 2")
            explorer.arena_combo.setCurrentIndex(index)
            explorer.arena_combo.activated.emit(index)
            self.assertEqual(panel._event_category(), "zone")
            self.assertTrue(explorer.isHidden())
            self.assertIn("Moving", panel._behavior_sources[fiber.stem]["behaviors"])
            self.assertIn("Zone: In zone", panel._behavior_sources[fiber.stem]["behaviors"])

    def test_analysis_kind_preserves_existing_zone_dashboard_and_windows(self):
        panel = self.panel
        time = np.arange(0., 30., .1)
        signal = np.sin(time)
        panel._processed = [ProcessedTrial(
            path="fiber.csv", channel_id="AIN01", time=time,
            raw_signal=signal, raw_reference=signal, output=signal, output_label="dFF")]
        panel._behavior_sources["fiber"] = _behavior_table_info(
            pd.DataFrame({"time": time, "In zone": ((time >= 12) & (time < 15)).astype(int)}),
            "binary_columns", 0.)
        panel._refresh_behavior_list()
        panel.spin_pre.setValue(3.)
        panel.spin_post.setValue(7.)
        panel.spin_metric_pre0.setValue(-2.)
        panel.spin_metric_post1.setValue(4.)
        panel._compute_psth()
        expected_matrix = panel._last_mat.copy()
        expected_events = panel._last_events.copy()
        old_plots = [panel._results_splitter.widget(i) for i in range(4)]
        self.assertTrue(panel.behavior_zone_panel.isHidden())
        panel._set_event_category("behavior")
        self.assertFalse(panel.behavior_zone_panel.isHidden())
        self.assertEqual(set(panel.behavior_zone_panel.selectors), {"behavior"})
        panel._set_event_category("zone")
        self.assertTrue(panel.behavior_zone_panel.isHidden())
        self.assertEqual([panel._results_splitter.widget(i) for i in range(4)], old_plots)
        self.assertEqual(panel.spin_pre.value(), 3.)
        self.assertEqual(panel.spin_post.value(), 7.)
        self.assertEqual(panel.spin_metric_pre0.value(), -2.)
        self.assertEqual(panel.spin_metric_post1.value(), 4.)
        self.assertFalse(hasattr(panel, "_analysis_tabs"))
        panel._compute_psth()
        np.testing.assert_allclose(panel._last_mat, expected_matrix, equal_nan=True)
        np.testing.assert_array_equal(panel._last_events, expected_events)

    def test_psth_behavior_onset_offset_transition_and_duration_filter(self):
        """MAMIR bouts use the same alignment and filtering controls as zones."""
        panel = self.panel
        time = np.arange(0., 40., .05)
        signal = np.sin(time)
        panel._processed = [ProcessedTrial(
            path="fiber.csv", channel_id="AIN01", time=time, raw_signal=signal,
            raw_reference=signal, output=signal, output_label="dFF")]
        panel._behavior_sources["fiber"] = {
            "kind": "binary_columns", "behaviors": {}, "trajectory": {},
            "event_behaviors": {
                "attack": {"on": np.array([10., 20.]), "off": np.array([12., 23.]),
                           "dur": np.array([2., 3.])},
                "grooming": {"on": np.array([12.5, 27.]), "off": np.array([13.5, 28.]),
                             "dur": np.array([1., 1.])},
                "Zone: target": {"on": np.array([5.]), "off": np.array([7.]),
                                 "dur": np.array([2.])},
            },
        }
        panel._refresh_behavior_list()
        for name, align, expected in (
            ("attack", "Align to onset", [10., 20.]),
            ("attack", "Align to offset", [12., 23.]),
            ("grooming", "Align to onset", [12.5, 27.]),
            ("Zone: target", "Align to onset", [5.]),
        ):
            panel.combo_behavior_name.setCurrentText(name)
            panel.combo_behavior_align.setCurrentText(align)
            panel._compute_psth()
            np.testing.assert_allclose(panel._last_events, expected)
            self.assertEqual(panel._last_mat.shape[0], len(expected))
        panel.combo_behavior_align.setCurrentText("Transition A->B")
        panel.combo_behavior_from.setCurrentText("attack")
        panel.combo_behavior_to.setCurrentText("grooming")
        panel.spin_transition_gap.setValue(1.)
        panel._compute_psth()
        np.testing.assert_allclose(panel._last_events, [12.5])
        panel.combo_behavior_name.setCurrentText("attack")
        panel.combo_behavior_align.setCurrentText("Align to onset")
        panel.spin_dur_min.setValue(2.5)
        panel._compute_psth()
        np.testing.assert_allclose(panel._last_events, [20.])

        panel._set_event_category("behavior")
        view = panel.behavior_zone_panel
        panel.combo_behavior_name.setCurrentText("grooming")
        panel.combo_behavior_name.activated.emit(panel.combo_behavior_name.currentIndex())
        self.assertEqual(view._selected_labels(), ["grooming"])
        attack = view.selectors["behavior"].findItems("attack", QtCore.Qt.MatchFlag.MatchExactly)[0]
        attack.setCheckState(QtCore.Qt.CheckState.Checked)
        panel.combo_behavior_name.setCurrentText("attack")
        panel.combo_behavior_name.activated.emit(panel.combo_behavior_name.currentIndex())
        self.assertEqual(set(view._selected_labels()), {"attack", "grooming"})

    def test_behavior_row_clicks_auto_update_and_coalesce_rapid_changes(self):
        """Text, checkboxes and keyboard update results without pressing Refresh."""
        panel = self.panel
        time = np.arange(0., 30., .1)
        for index in range(2):
            path = f"fiber{index}.csv"
            signal = time + index
            panel._processed.append(ProcessedTrial(
                path=path, channel_id="AIN01", time=time, raw_signal=signal,
                raw_reference=signal, output=signal, output_label="dFF"))
            panel._behavior_sources[Path(path).stem] = _behavior_table_info(
                pd.DataFrame({"time": time, "attack": ((time >= 10) & (time < 12)).astype(int),
                              "grooming": ((time >= 20) & (time < 23)).astype(int)}),
                "binary_columns", 0.)
        panel._refresh_behavior_list()
        panel._set_event_category("behavior")
        panel.resize(1400, 1000)
        panel.show()
        QtTest.QTest.qWait(450)
        view = panel.behavior_zone_panel
        listing = view.selectors["behavior"]
        groom = listing.findItems("grooming", QtCore.Qt.MatchFlag.MatchExactly)[0]
        point = listing.visualItemRect(groom).center()
        with patch.object(view, "_compute", wraps=view._compute) as compute:
            for _ in range(3):
                QtTest.QTest.mouseClick(listing.viewport(), QtCore.Qt.MouseButton.LeftButton, pos=point)
            self.assertEqual(set(view._selected_labels()), {"attack", "grooming"})
            self.assertEqual(panel.combo_behavior_name.currentText(), "grooming")
            QtTest.QTest.qWait(450)
            self.assertEqual(compute.call_count, 1)
            self.assertEqual(view.table.rowCount(), 2)
            self.assertTrue(view.export_button.isEnabled())
        # A click on the checkbox must toggle exactly once too.
        rect = listing.visualItemRect(groom)
        QtTest.QTest.mouseClick(listing.viewport(), QtCore.Qt.MouseButton.LeftButton,
                               pos=QtCore.QPoint(rect.left() + 8, rect.center().y()))
        QtTest.QTest.qWait(300)
        self.assertEqual(view._selected_labels(), ["attack"])
        self.assertEqual(panel.combo_behavior_name.currentText(), "attack")
        self.assertEqual(view.table.rowCount(), 1)
        QtTest.QTest.keyClick(listing, QtCore.Qt.Key.Key_Space)
        QtTest.QTest.qWait(300)
        self.assertEqual(view.table.rowCount(), 2)
        with patch.object(view, "_compute", wraps=view._compute) as compute:
            view.before.setValue(3.)
            view.offset.setValue(.5)
            panel.tab_visual_mode.setCurrentIndex(1)
            QtTest.QTest.qWait(450)
            self.assertEqual(compute.call_count, 1)
            self.assertEqual(len(view._last_group_rows), 2)
            self.assertTrue(all(row["recordings"] == 2 for row in view._last_group_rows))
        for index in range(listing.count()):
            listing.item(index).setCheckState(QtCore.Qt.CheckState.Unchecked)
        QtTest.QTest.qWait(300)
        self.assertEqual(view.table.rowCount(), 0)
        self.assertFalse(view.export_button.isEnabled())
        self.assertFalse(view._compare_timer.isActive())
        view.reload()
        self.assertEqual(view._selected_labels(), [])

    def test_behavior_timing_help_follows_psth_without_changing_zone_windows(self):
        panel = self.panel
        view = panel.behavior_zone_panel
        panel.spin_pre.setValue(4.)
        panel.spin_post.setValue(7.)
        panel._set_event_category("behavior")
        view.before.setValue(2.)
        view.after.setValue(3.)
        panel.combo_behavior_align.setCurrentText("Align to offset")
        self.assertIn("2 s before start", view.timing_help.text())
        self.assertIn("3 s after end", view.timing_help.text())
        self.assertIn("0 = behavior end", view.timing_help.text())
        self.assertIn("−4 to +7 s", view.timing_help.text())
        view.offset.setValue(.5)
        self.assertIn("heatmap is NOT shifted", view.timing_help.text())
        panel._set_event_category("zone")
        self.assertEqual((panel.spin_pre.value(), panel.spin_post.value()), (4., 7.))
        self.assertTrue(view.isHidden())

    def test_large_behavior_targets_filter_and_bulk_selection(self):
        panel = self.panel
        panel._behavior_sources = {"fiber": _behavior_table_info(pd.DataFrame({
            "time": [0., 1., 2.], "attack": [0, 1, 0], "grooming": [0, 1, 0],
            "approach": [0, 0, 1]}), "binary_columns", 0.)}
        panel._refresh_behavior_list()
        panel._set_event_category("behavior")
        panel.resize(1300, 1000)
        panel.show()
        self.app.processEvents()
        view = panel.behavior_zone_panel
        listing = view.selectors["behavior"]
        view.clear_selection_button.click()
        item = listing.findItems("grooming", QtCore.Qt.MatchFlag.MatchExactly)[0]
        rect = listing.visualItemRect(item)
        self.assertGreaterEqual(rect.width(), 210)
        self.assertGreaterEqual(rect.height(), 38)
        QtTest.QTest.mouseClick(listing.viewport(), QtCore.Qt.MouseButton.LeftButton,
                               pos=QtCore.QPoint(rect.right() - 10, rect.center().y()))
        self.assertEqual(view._selected_labels(), ["grooming"])
        view.searches["behavior"].setText("att")
        view.select_visible_button.click()
        self.assertEqual(set(view._selected_labels()), {"attack", "grooming"})
        view.clear_selection_button.click()
        self.assertEqual(view._selected_labels(), [])
        view.searches["behavior"].clear()
        view.reload()
        self.assertEqual(view._selected_labels(), [])

    def test_behavior_psth_controls_and_resize_preserve_analysis(self):
        panel = self.panel
        time = np.arange(0., 35., .1)
        panel._processed = [ProcessedTrial(path="fiber.csv", channel_id="AIN01", time=time,
                                           raw_signal=time, raw_reference=time,
                                           output=np.sin(time), output_label="dFF")]
        panel._behavior_sources = {"fiber": _behavior_table_info(pd.DataFrame({
            "time": time, "attack": ((time >= 10) & (time < 12)).astype(int),
            "grooming": ((time >= 20) & (time < 23)).astype(int)}), "binary_columns", 0.)}
        panel._refresh_behavior_list()
        panel._set_event_category("behavior")
        panel.resize(1400, 1000)
        panel.show()
        self.app.processEvents()
        bar = panel.behavior_psth_bar
        bar.name.setCurrentText("grooming")
        bar.name.activated.emit(bar.name.currentIndex())
        self.assertEqual(panel.combo_behavior_name.currentText(), "grooming")
        bar.pre.setValue(3.)
        bar.post.setValue(6.)
        bar.align.setCurrentText("Align to offset")
        bar.align.activated.emit(bar.align.currentIndex())
        self.assertEqual((panel.spin_pre.value(), panel.spin_post.value()), (3., 6.))
        self.assertEqual(panel.combo_behavior_align.currentText(), "Align to offset")
        panel.spin_post.setValue(4.)
        self.assertEqual(bar.post.value(), 4.)
        panel._compute_psth()
        expected = panel._last_mat.copy()
        self.assertIn("grooming", panel.plot_heat.plotItem.titleLabel.text)
        self.assertIn("behavior end", bar.summary.text())
        self.assertIn("individual bouts", bar.summary.text())
        sizes = panel._behavior_workspace.sizes()
        panel._behavior_workspace.moveSplitter(sizes[0] + 60, 1)
        self.app.processEvents()
        resized = panel._behavior_workspace.sizes()
        self.assertNotEqual(sizes, resized)
        self.assertEqual(panel._behavior_workspace_sizes, resized)
        self.assertEqual([int(value) for value in panel._settings.value("behavior_workspace_sizes")], resized)
        bar.collapse_button.click()
        self.assertTrue(panel._behavior_comparison_scroll.isHidden())
        self.assertFalse(bar.isHidden())
        bar.collapse_button.click()
        self.assertFalse(panel._behavior_comparison_scroll.isHidden())
        np.testing.assert_equal(panel._last_mat, expected)
        bar.heat_button.click()
        self.assertGreater(panel._results_scroll.verticalScrollBar().value(), 0)
        bar.mean_button.click()
        np.testing.assert_equal(panel._last_mat, expected)
        panel._set_event_category("zone")
        self.assertTrue(bar.isHidden())
        self.assertTrue(panel._behavior_comparison_scroll.isHidden())
        self.assertEqual(panel.plot_heat.plotItem.titleLabel.text, "Heatmap")
        self.assertEqual(panel.plot_heat.getAxis("bottom").labelText, "")
        self.assertFalse(panel.plot_heat.getAxis("bottom").style["showValues"])
        np.testing.assert_equal(panel._last_mat, expected)

    def test_behavior_group_rejects_mixed_signal_units(self):
        panel = self.panel
        time = np.arange(0., 20., .1)
        for index, units in enumerate(("dFF", "z-score")):
            panel._processed.append(ProcessedTrial(
                path=f"fiber_{index}.csv", channel_id="AIN01", time=time,
                raw_signal=time, raw_reference=time, output=time, output_label=units))
            panel._behavior_sources[f"fiber_{index}"] = _behavior_table_info(
                pd.DataFrame({"time": time, "attack": ((time >= 5) & (time < 7)).astype(int)}),
                "binary_columns", 0.)
        panel._refresh_behavior_list()
        panel._set_event_category("behavior")
        panel.tab_visual_mode.setCurrentIndex(1)
        view = panel.behavior_zone_panel
        view._compute()
        self.assertIn("different signal units", view.status.text())
        self.assertEqual(view.table.rowCount(), 0)
        self.assertFalse(view.export_button.isEnabled())

    def test_metadata_batch_matches_recordings_before_load_order(self):
        """Reversed metadata import order cannot swap the recording association."""
        panel = self.panel
        first = SimpleNamespace(path="29539_1_baseline.csv")
        second = SimpleNamespace(path="29539_2_baseline.csv")
        panel._processed = [first, second]
        one, two = {"recording": 1}, {"recording": 2}
        panel._behavior_sources = {"29539_2_baseline_metadata": two,
                                   "29539_1_baseline_metadata": one}
        self.assertIs(panel._match_behavior_source(first), one)
        self.assertIs(panel._match_behavior_source(second), two)
        panel._processed = []

    def test_comparison_settings_survive_project_roundtrip_and_clear_selection(self):
        panel = self.panel
        time = np.arange(0., 30., .1)
        panel._processed = [ProcessedTrial(path="fiber.csv", channel_id="AIN01", time=time,
                                           raw_signal=time, raw_reference=time, output=time, output_label="dFF")]
        panel._behavior_sources = {"fiber": _behavior_table_info(pd.DataFrame({
            "time": time, "attack": ((time > 10) & (time < 12)).astype(int),
            "grooming": ((time > 20) & (time < 23)).astype(int)}), "binary_columns", 0.)}
        panel._refresh_behavior_list()
        panel._set_event_category("behavior")
        view = panel.behavior_zone_panel
        view._set_visible_selection(True)
        view.before.setValue(3.5)
        view.after.setValue(4.5)
        view.offset.setValue(.25)
        view.searches["behavior"].setText("attack")
        panel.behavior_psth_bar.collapse_button.setChecked(True)
        expected = view.settings_state()
        with tempfile.TemporaryDirectory() as directory:
            path = str(Path(directory) / "comparison.h5")
            panel._save_project_h5(path)
            view.before.setValue(1.)
            view.clear_selection_button.click()
            view.searches["behavior"].clear()
            panel.behavior_psth_bar.collapse_button.setChecked(False)
            self.assertTrue(panel._load_project_from_path(path))
            self.assertEqual(view.settings_state(), expected)
            view.clear_selection_button.click()
            panel._save_project_h5(path)
            view._set_visible_selection(True)
            panel._load_project_from_path(path)
            self.assertEqual(view._selected_labels(), [])
            view.reload()
            self.assertEqual(view._selected_labels(), [])

    def test_recent_mamir_file_uses_auto_detection_and_keeps_other_sources(self):
        panel = self.panel
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "behavior_frames.csv"
            pd.DataFrame({"frame": [0, 1, 2], "time_s": [0., .1, .2], "identity": [1, 1, 1],
                          "walking": [1, 0, 0], "attack": [0, 1, 0]}).to_csv(path, index=False)
            prior = _behavior_table_info(pd.DataFrame({"time": [0, 1], "legacy": [0, 1]}),
                                        "binary_columns", 0.)
            panel._behavior_sources["untouched"] = prior
            panel._load_recent_behavior_path(str(path))
            self.assertIs(panel._behavior_sources["untouched"], prior)
            events = panel._behavior_sources[path.stem]["event_behaviors"]
            np.testing.assert_allclose(events["walking"]["on"], [0.])
            np.testing.assert_allclose(events["attack"]["on"], [.1])


if __name__ == "__main__":
    unittest.main()
